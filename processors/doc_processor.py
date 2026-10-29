import io
from typing import Dict, Optional, Tuple
from docx import Document
from engine.humanizer import TurnitoutEngine


def process_docx_in_place(
    file_bytes: bytes,
    engine: TurnitoutEngine,
    intensity: int = 85,
    register: str = "academic",
    use_zipf_synsets: bool = True,
    voice_profile: Optional[Dict[str, float]] = None
) -> Tuple[bytes, Dict[str, int]]:
    doc = Document(io.BytesIO(file_bytes))
    agg_telemetry = {
        "watermarks_purged": 0,
        "ai_tells_scrubbed": 0,
        "discourse_rotated": 0,
        "register_shifts": 0,
        "zipf_inflections": 0,
        "clauses_coordinated": 0,
        "firewall_intercepts": 0,
        "invariants_protected": 0
    }

    cp = doc.core_properties
    cp.author = ""
    cp.last_modified_by = ""
    cp.comments = ""
    cp.subject = ""
    cp.keywords = ""
    cp.category = ""
    cp.content_status = ""
    cp.revision = 1

    def _process_paragraphs(paragraphs):
        for para in paragraphs:
            style_name = (para.style.name if para.style else "").lower()
            if any(skip in style_name for skip in ["code", "macro", "toc", "bibliography", "heading", "title", "caption"]):
                continue

            full_text = para.text
            if not full_text or len(full_text.strip().split()) < 5:
                continue

            leading_ws = len(full_text) - len(full_text.lstrip(" "))
            trailing_ws = len(full_text) - len(full_text.rstrip(" "))
            core_str = full_text.strip(" ")

            transformed, tel = engine.humanize_text(
                core_str,
                intensity=intensity,
                register=register,
                use_zipf_synsets=use_zipf_synsets,
                voice_profile=voice_profile
            )
            final_str = (" " * leading_ws) + transformed + (" " * trailing_ws)

            # Assign to first run to preserve paragraph font/size styling without XML run fragmentation
            if para.runs:
                para.runs[0].text = final_str
                for extra_run in para.runs[1:]:
                    extra_run.text = ""
            else:
                para.text = final_str

            for k in agg_telemetry:
                agg_telemetry[k] += tel.get(k, 0)

    _process_paragraphs(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                _process_paragraphs(cell.paragraphs)

    out_buf = io.BytesIO()
    doc.save(out_buf)
    return out_buf.getvalue(), agg_telemetry